from docx import Document
from docx.shared import Cm

from app.document import build_docx
from app.planner import normalise_plan
from app.schemas import Block, QualityReport, Section
from app.tools import ToolRegistry
from tests.fakes import GOOD_PLAN


def _build(tmp_path):
    plan = normalise_plan(GOOD_PLAN, "req", ToolRegistry())
    secs = [Section(task_id=t.id, title=t.section_title, key_points=[], blocks=[
        Block(type="paragraph", text=f"Body of {t.id}"), Block(type="bullets", items=["a", "b"]),
        Block(type="numbered", items=["x", "y"]), Block(type="table", caption="Cap", headers=["H1", "H2"], rows=[["1", "2"], ["3", "4"]])],
        tool_log=[{"tool": "calculate", "args": {"expression": "1+1"}, "result": {"value": 2}}] if t.id == "T2" else []) for t in plan.tasks]
    name, path = build_docx(request_id="abcdef123456", original_request="req", plan=plan, sections=secs, quality=QualityReport(status="passed"), warnings=["note"], output_dir=tmp_path)
    return name, path, Document(path)


def test_document_structure_and_appendices(tmp_path):
    name, path, doc = _build(tmp_path)
    assert name.endswith("abcdef12.docx") and path.stat().st_size > 5000
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert "1. Executive Summary" in headings and "Appendix A: Agent execution plan" in headings
    assert "Appendix B: Requirements traceability" in headings and "Appendix C: Tool calls" in headings
    assert len(doc.tables) >= 3 + 4  # title block + 4 section tables + appendices
    assert abs(doc.sections[0].page_width - Cm(21.0)) < 1000 and abs(doc.sections[0].page_height - Cm(29.7)) < 1000  # A4


def test_header_repeat_and_page_number_fields(tmp_path):
    _, path, doc = _build(tmp_path)
    xml = doc.sections[0].footer._element.xml
    assert "PAGE" in xml and "NUMPAGES" in xml
    assert "w:tblHeader" in doc.tables[1]._tbl.xml
