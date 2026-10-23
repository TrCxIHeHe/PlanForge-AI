import json
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document

from app.agent import validate_request
from app.config import Settings

EXAMPLES = Path(__file__).resolve().parents[1] / "examples"


def test_example_requests_pass_guardrails():
    for f in EXAMPLES.glob("request_*.json"):
        text = json.loads(f.read_text(encoding="utf-8"))["request"]
        assert validate_request(text, Settings())



def test_demo_script_downloads_document_when_base64_is_disabled(tmp_path, monkeypatch):
    import sys
    import httpx
    from scripts import run_examples

    examples = tmp_path / "examples"
    examples.mkdir()
    (examples / "request_standard.json").write_text(
        json.dumps({"request": "Prepare a useful standard business proposal."}), encoding="utf-8"
    )
    run_examples.ROOT = tmp_path
    payload = {
        "summary": "Created example.",
        "plan": {
            "title": "Example Proposal",
            "source": "llm",
            "repairs": [],
            "assumptions": [],
            "tasks": [{"section_title": "Project Approach"}],
        },
        "execution": [],
        "quality_check": {"status": "passed", "revised_tasks": []},
        "llm": {"provider": "fake", "calls_by_stage": {}},
        "warnings": [],
        "document": {"base64_content": None, "download_path": "/documents/example.docx"},
    }
    post_request = httpx.Request("POST", "http://test/agent")
    get_request = httpx.Request("GET", "http://test/documents/example.docx")
    document = Document()
    document.add_heading("Example Proposal", 0)
    document.add_heading("Project Approach", 1)
    document.add_paragraph("This is substantive content for the validated Word document. " * 3)
    buffer = BytesIO()
    document.save(buffer)
    valid_docx = buffer.getvalue()
    monkeypatch.setattr(run_examples.httpx, "post", lambda *a, **k: httpx.Response(200, json=payload, request=post_request))
    monkeypatch.setattr(run_examples.httpx, "get", lambda *a, **k: httpx.Response(200, content=valid_docx, request=get_request))
    monkeypatch.setattr(sys, "argv", ["run_examples.py", "standard", "--url", "http://test"])

    assert run_examples.main() == 0
    saved = tmp_path / "output" / "demo" / "standard.docx"
    assert Document(saved).paragraphs[0].text == "Example Proposal"


def test_demo_script_rejects_malformed_pk_document_without_saving(tmp_path, monkeypatch):
    import sys
    import httpx
    from scripts import run_examples

    examples = tmp_path / "examples"
    examples.mkdir()
    (examples / "request_standard.json").write_text(
        json.dumps({"request": "Prepare a useful standard business proposal."}), encoding="utf-8"
    )
    run_examples.ROOT = tmp_path
    payload = {
        "summary": "Created example.",
        "plan": {
            "title": "Example Proposal",
            "source": "llm",
            "repairs": [],
            "assumptions": [],
            "tasks": [{"section_title": "Project Approach"}],
        },
        "execution": [],
        "quality_check": {"status": "passed", "revised_tasks": []},
        "llm": {"provider": "fake", "calls_by_stage": {}},
        "warnings": [],
        "document": {"base64_content": None, "download_path": "/documents/example.docx"},
    }
    post_request = httpx.Request("POST", "http://test/agent")
    get_request = httpx.Request("GET", "http://test/documents/example.docx")
    monkeypatch.setattr(run_examples.httpx, "post", lambda *a, **k: httpx.Response(200, json=payload, request=post_request))
    monkeypatch.setattr(
        run_examples.httpx, "get",
        lambda *a, **k: httpx.Response(200, content=b"PK\x03\x04malformed", request=get_request),
    )
    monkeypatch.setattr(sys, "argv", ["run_examples.py", "standard", "--url", "http://test"])

    with pytest.raises(RuntimeError, match="not a valid \\.docx"):
        run_examples.main()
    assert not (tmp_path / "output" / "demo" / "standard.docx").exists()
