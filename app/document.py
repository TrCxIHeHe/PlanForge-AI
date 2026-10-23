"""Professional Microsoft Word output: A4, styled tables, header/footer with page numbers, plan appendix."""
from __future__ import annotations

import json
import re
from datetime import datetime
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .schemas import AgentPlan, Block, QualityReport, Section

NAVY, BLUE, GREY = RGBColor(31, 55, 82), RGBColor(47, 84, 150), RGBColor(95, 105, 115)
FONT = "Calibri"


def _slug(text: str) -> str:
    return re.sub(r"[^a-zA-Z0-9]+", "-", text).strip("-").lower()[:48] or "agent-document"


def _shade(cell, fill: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    props.append(shd)


def _field(paragraph, instruction: str, size: float = 8) -> None:
    for kind, text in (("begin", None), (None, instruction), ("separate", None), (None, "1"), ("end", None)):
        run = paragraph.add_run()
        run.font.size, run.font.color.rgb = Pt(size), GREY
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        elif text == instruction:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = f" {instruction} "
        else:
            el = OxmlElement("w:t")
            el.text = text
        run._r.append(el)


def _repeat_header(row) -> None:
    pr = row._tr.get_or_add_trPr()
    for tag in ("w:tblHeader", "w:cantSplit"):
        pr.append(OxmlElement(tag))


def _table(doc, headers: list[str], rows: Iterable[list[str]], caption: str = "", widths: list[float] | None = None) -> None:
    rows = list(rows)
    if caption:
        p = doc.add_paragraph()
        r = p.add_run(caption)
        r.bold, r.font.size, r.font.color.rgb = True, Pt(9), GREY
        p.paragraph_format.space_after, p.paragraph_format.keep_with_next = Pt(3), True
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        run.bold, run.font.size, run.font.color.rgb = True, Pt(9), RGBColor(255, 255, 255)
        _shade(cell, "1F3752")
    _repeat_header(table.rows[0])
    for n, row in enumerate(rows):
        cells = table.add_row().cells
        table.rows[-1]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for i, value in enumerate(row[: len(headers)]):
            cells[i].text = ""
            para = cells[i].paragraphs[0]
            para.paragraph_format.space_after = Pt(2)
            run = para.add_run(str(value))
            run.font.size = Pt(9)
            if n % 2:
                _shade(cells[i], "F2F5F9")
    if not widths:  # equal columns, but never wider than the text area
        widths = [16.6 / len(headers)] * len(headers)
    table.autofit = False
    for col, w in zip(table.columns, widths):
        col.width = Cm(w)
    for row in table.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def _render_block(doc, block: Block) -> None:
    if block.type == "paragraph":
        doc.add_paragraph(block.text)
    elif block.type in ("bullets", "numbered"):
        style = "List Bullet" if block.type == "bullets" else "List Number"
        for item in block.items:
            p = doc.add_paragraph(item, style=style)
            p.paragraph_format.space_after = Pt(2)
    else:
        _table(doc, block.headers, block.rows, block.caption)


def build_docx(
    *,
    request_id: str,
    original_request: str,
    plan: AgentPlan,
    sections: list[Section],
    quality: QualityReport,
    warnings: list[str],
    output_dir: Path,
) -> tuple[str, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    filename = f"{_slug(plan.title)}-{request_id[:8]}.docx"
    path = output_dir / filename

    doc = Document()
    page = doc.sections[0]
    page.page_width, page.page_height = Cm(21.0), Cm(29.7)  # A4
    page.top_margin = page.bottom_margin = Cm(2.0)
    page.left_margin = page.right_margin = Cm(2.2)

    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = FONT, Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.color.rgb = RGBColor(43, 48, 53)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12
    for name, size, color in (("Title", 26, NAVY), ("Heading 1", 15, NAVY), ("Heading 2", 12, BLUE)):
        st = doc.styles[name]
        st.font.name, st.font.size, st.font.bold, st.font.color.rgb = FONT, Pt(size), True, color
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(16)
    doc.styles["Heading 1"].paragraph_format.keep_with_next = True

    # header / footer
    hp = page.header.paragraphs[0]
    hr = hp.add_run(plan.title)
    hr.font.size, hr.font.color.rgb = Pt(8), GREY
    fp = page.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = fp.add_run("Page ")
    r.font.size, r.font.color.rgb = Pt(8), GREY
    _field(fp, "PAGE")
    r = fp.add_run(" of ")
    r.font.size, r.font.color.rgb = Pt(8), GREY
    _field(fp, "NUMPAGES")

    # title block
    doc.add_paragraph(plan.title, style="Title")
    sub = doc.add_paragraph()
    run = sub.add_run(plan.document_type.title())
    run.font.size, run.font.color.rgb, run.bold = Pt(12), BLUE, True
    _table(
        doc, ["Item", "Detail"],
        [["Prepared", datetime.now().strftime("%d %B %Y")], ["Generated by", "Autonomous Document Agent"], ["Reference", request_id[:8]]],
        widths=[4.0, 12.6],
    )

    doc.add_heading("Purpose", level=1)
    doc.add_paragraph(plan.purpose)
    if plan.assumptions:
        doc.add_heading("Assumptions", level=2)
        for a in plan.assumptions:
            doc.add_paragraph(a, style="List Bullet").paragraph_format.space_after = Pt(2)
    if warnings:
        doc.add_heading("Review notes", level=2)
        for w in warnings:
            doc.add_paragraph(w, style="List Bullet").paragraph_format.space_after = Pt(2)

    for n, section in enumerate(sections, 1):
        doc.add_heading(f"{n}. {section.title}", level=1)
        for block in section.blocks:
            _render_block(doc, block)

    # appendix: how the agent produced the document
    doc.add_page_break()
    doc.add_heading("Appendix A: Agent execution plan", level=1)
    doc.add_paragraph(f"Plan source: {plan.source}. Tasks with no dependencies ran in parallel; each task is traced to the requirements it satisfies.")
    status = {s.task_id: s for s in sections}
    _table(
        doc, ["ID", "Task", "Depends on", "Requirements", "Tools", "Status"],
        [[t.id, t.title, ", ".join(t.depends_on) or "-", ", ".join(t.satisfies) or "-",
          ", ".join(sorted({c["tool"] for c in status[t.id].tool_log if "result" in c})) or "-", status[t.id].status] for t in plan.tasks],
        widths=[1.1, 5.6, 2.4, 2.6, 3.0, 2.0],
    )
    doc.add_heading("Appendix B: Requirements traceability", level=1)
    owners = {r.id: [t.id for t in plan.tasks if r.id in t.satisfies] for r in plan.requirements}
    checks = {c.requirement_id: c for c in quality.checks}
    _table(
        doc, ["ID", "Requirement", "Covered by", "Review"],
        [[r.id, r.text, ", ".join(owners[r.id]), (checks[r.id].status.replace("_", " ") + (f": {checks[r.id].note}" if checks[r.id].note else "")) if r.id in checks else "-"] for r in plan.requirements],
        widths=[1.1, 8.0, 2.4, 5.2],
    )
    if quality.summary or quality.revised_tasks:
        text = quality.summary
        if quality.revised_tasks:
            text += f" Sections revised after self-review: {', '.join(quality.revised_tasks)}."
        doc.add_paragraph(text.strip())

    tool_rows = []
    for s in sections:
        for c in s.tool_log:
            body = json.dumps(c.get("result", c.get("error")), default=str)
            tool_rows.append([s.task_id, c["tool"], json.dumps(c.get("args"), default=str)[:140], body[:200] + ("..." if len(body) > 200 else "")])
    if tool_rows:
        doc.add_heading("Appendix C: Tool calls", level=1)
        _table(doc, ["Task", "Tool", "Arguments", "Result"], tool_rows, widths=[1.3, 2.4, 6.0, 7.0])

    doc.core_properties.title, doc.core_properties.author = plan.title, "Autonomous Document Agent"
    doc.save(path)
    return filename, path
