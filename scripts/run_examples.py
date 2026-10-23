"""Send the two example requests to a running server and save the generated Word files.

Usage: python scripts/run_examples.py [standard|complex|all] [--url http://127.0.0.1:8000]
"""
import argparse
import base64
from io import BytesIO
import json
import sys
from pathlib import Path

import httpx
from docx import Document

ROOT = Path(__file__).resolve().parents[1]


def validate_docx(content: bytes, *, title: str, section_titles: list[str]) -> None:
    """Require a readable Word package containing the expected generated content."""
    try:
        document = Document(BytesIO(content))
    except Exception as exc:
        raise RuntimeError("Downloaded document is not a valid .docx file.") from exc

    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    paragraphs.extend(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    text = "\n".join(paragraphs)
    folded = text.casefold()
    if len(text.strip()) < 80:
        raise RuntimeError("Downloaded .docx contains no substantive document content.")
    if title and title.casefold() not in folded:
        raise RuntimeError("Downloaded .docx is missing the expected document title.")
    missing_sections = [section for section in section_titles if section.casefold() not in folded]
    if missing_sections:
        raise RuntimeError(
            "Downloaded .docx is missing expected section(s): " + ", ".join(missing_sections)
        )


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("which", nargs="?", default="all", choices=["standard", "complex", "all"])
    ap.add_argument("--url", default="http://127.0.0.1:8000")
    args = ap.parse_args()
    names = ["standard", "complex"] if args.which == "all" else [args.which]
    out_dir = ROOT / "output" / "demo"
    out_dir.mkdir(parents=True, exist_ok=True)
    for name in names:
        payload = json.loads((ROOT / "examples" / f"request_{name}.json").read_text(encoding="utf-8"))
        print(f"\n=== {name} request ===")
        r = httpx.post(f"{args.url}/agent", json=payload, timeout=1800)
        if r.status_code != 200:
            print("FAILED", r.status_code, r.text)
            return 1
        body = r.json()
        print(body["summary"])
        print("plan source:", body["plan"]["source"], "| repairs:", body["plan"]["repairs"] or "none")
        print("assumptions:", *body["plan"]["assumptions"], sep="\n  - ")
        print("task list:")
        for t in body["execution"]:
            print(f"  {t['task_id']} [{t['status']}] {t['title']}  deps={t['depends_on']} tools={t['tools_used']}")
        print("quality:", body["quality_check"]["status"], "| revised:", body["quality_check"]["revised_tasks"] or "none")
        print("llm:", {k: v for k, v in body["llm"].items() if k != "calls_by_stage"})
        for w in body["warnings"]:
            print("WARNING:", w)
        path = out_dir / f"{name}.docx"
        encoded = body["document"].get("base64_content")
        if encoded:
            try:
                content = base64.b64decode(encoded, validate=True)
            except (ValueError, TypeError) as exc:
                raise RuntimeError(f"Invalid base64 document returned for {name}.") from exc
        else:  # INCLUDE_DOCUMENT_BASE64=false: fetch through the download endpoint instead
            try:
                download = httpx.get(args.url + body["document"]["download_path"], timeout=60)
                download.raise_for_status()
            except httpx.HTTPError as exc:
                raise RuntimeError(f"Document download failed for {name}: {exc}") from exc
            content = download.content
        validate_docx(
            content,
            title=str(body.get("plan", {}).get("title") or ""),
            section_titles=[
                str(task.get("section_title") or "")
                for task in body.get("plan", {}).get("tasks", [])
                if task.get("section_title")
            ],
        )
        path.write_bytes(content)
        print("saved", path)
    return 0


if __name__ == "__main__":
    sys.exit(main())
